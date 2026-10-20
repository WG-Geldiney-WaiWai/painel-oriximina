"""Prepara o painel "Planejamento urbano" de Oriximiná (como o do Taipei, que mostra o zoneamento da cidade).

O zoneamento do Plano Diretor de Oriximiná não está publicado como dados; os cartões mostram o que existe
para o município inteiro:

- comp_310.csv: uso e cobertura da terra em 2025, em km² (MapBiomas, Coleção 11)
- comp_310_historico.csv: como pastagem, mineração, área urbana e agricultura cresceram desde 1985
- comp_311.csv: áreas protegidas dentro de Oriximiná, em km²: terras indígenas (FUNAI), territórios
  quilombolas (INCRA) e unidades de conservação (MMA, CNUC). As áreas se sobrepõem em vários lugares; o total
  sem contar duas vezes vai na descrição

Os arquivos baixados ficam na pasta das fontes, em MapBiomas e AreasProtegidas (são usados de novo se já
estiverem lá; apague-os para baixar versões novas). O contorno do município vem de PA_Municipios_2025.

Rode de novo para atualizar:  python preparar_planejamento.py
"""
import glob
import json
import os
import re
import urllib.request
import warnings
import zipfile

import geopandas as gpd
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")

PASTA_PROJETO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
PASTA_COMPONENTES = os.path.join(PASTA_PROJETO, "dados", "componentes")
ARQ_ADMIN = os.path.join(PASTA_PROJETO, "usuarios", "admin_painel.json")
PASTA_FONTES = os.environ.get("PASTA_FONTES", r"C:\Users\Usuário\Downloads\Pasta Mapa de Painel Oriximiná")
CODIGO_IBGE = "1505304"
LIMITES = (-58.95, -3.4, -54.8, 2.6)   # retângulo em volta de Oriximiná (lon/lat), para baixar só a região
# Projeção de área igual (Albers para a América do Sul): os km² saem certos
ALBERS = "+proj=aea +lat_1=-5 +lat_2=-42 +lat_0=-32 +lon_0=-60 +ellps=GRS80 +units=m"

# MapBiomas Brasil, Coleção 11: área por classe de cobertura em cada município (1985 a 2025)
ID_MAPBIOMAS = "1otOqymHuixvkRGVl65zTTNyfaHo46Gqk"   # no Google Drive, link da página de estatísticas
URL_TIS = ("https://geoserver.funai.gov.br/geoserver/Funai/ows?service=WFS&version=1.0.0&request=GetFeature"
           "&typeName=Funai:tis_poligonais&outputFormat=application/json&bbox={},{},{},{}".format(*LIMITES))
URL_QUILOMBOLAS = "https://certificacao.incra.gov.br/csv_shp/zip/%C3%81reas%20de%20Quilombolas.zip"
API_CNUC = "https://dados.mma.gov.br/api/3/action/package_show?id=unidadesdeconservacao"

# Classes do MapBiomas (código: nome no cartão)
CLASSES = {
    3: "Floresta", 6: "Floresta alagável", 4: "Formação savânica", 12: "Campo", 11: "Campo alagado e área úmida",
    29: "Afloramento rochoso", 33: "Rios e lagos", 15: "Pastagem", 41: "Agricultura", 21: "Mosaico de usos",
    30: "Mineração", 24: "Área urbana", 25: "Outras áreas sem vegetação",
}
HISTORICO = [15, 30, 24, 41]   # pastagem, mineração, área urbana, agricultura: o que a ocupação humana mudou

ACENTOS = {"Grao": "Grão", "Para": "Pará", "Biologica": "Biológica", "Ecologica": "Ecológica", "Estacao": "Estação",
           "Saraca": "Saracá", "Abui": "Abuí", "Parana": "Paraná", "Coracao": "Coração", "Agua": "Água", "Zoe": "Zo'é"}


def baixar(url, destino, dados=None):
    if not os.path.exists(destino):
        print("baixando", url, flush=True)
        req = urllib.request.Request(url, data=dados, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=900) as resposta, open(destino, "wb") as arquivo:
            arquivo.write(resposta.read())
    return destino


def baixar_do_drive(id_arquivo, destino):
    """O Google Drive pede uma confirmação (aviso de antivírus) antes dos arquivos grandes."""
    if os.path.exists(destino):
        return destino
    pagina = urllib.request.urlopen(urllib.request.Request(
        f"https://drive.google.com/uc?id={id_arquivo}&export=download", headers={"User-Agent": "Mozilla/5.0"}),
        timeout=120).read()
    uuid = re.search(rb'name="uuid" value="([^"]+)"', pagina)
    url = (f"https://drive.usercontent.google.com/download?id={id_arquivo}&export=download&confirm=t"
           + (f"&uuid={uuid.group(1).decode()}" if uuid else ""))
    return baixar(url, destino)


def sem_mojibake(texto):
    """O CNUC vem em UTF-8 com campos cortados no meio de uma letra (a leitura em UTF-8 falha): é lido em
    latin1, que guarda os bytes como estão, e volta para UTF-8 aqui (a letra cortada no fim se perde)."""
    if not isinstance(texto, str):
        return texto
    return texto.encode("latin1").decode("utf-8", errors="ignore")


def nome_bonito(nome):
    nome = re.sub(r"\s+", " ", str(nome)).strip().title()
    nome = re.sub(r"\b(De|Da|Do|Dos|Das|E)\b", lambda m: m.group(1).lower(), nome)
    nome = re.sub(r"\b(Ii|Iii|Iv|Tq)\b", lambda m: m.group(1).upper(), nome)
    return re.sub(r"\b\w+\b", lambda m: ACENTOS.get(m.group(0), m.group(0)), nome).replace(",", " –")


def km2(n):
    return f"{n:,.0f}".replace(",", ".")


def cobertura():
    pasta = os.path.join(PASTA_FONTES, "MapBiomas")
    os.makedirs(pasta, exist_ok=True)
    arquivo = baixar_do_drive(ID_MAPBIOMAS, os.path.join(pasta, "mapbiomas_col11_municipios.zip"))
    if zipfile.is_zipfile(arquivo) and not arquivo.endswith(".xlsx"):
        with zipfile.ZipFile(arquivo) as z:
            nome = next(n for n in z.namelist() if n.endswith(".xlsx"))
            if not os.path.exists(os.path.join(pasta, nome)):
                z.extract(nome, pasta)
        arquivo = os.path.join(pasta, nome)
    print("lendo", arquivo, "(uns 30 segundos)", flush=True)
    # Linha a linha (a planilha tem o Brasil inteiro): fica só com as linhas de Oriximiná
    planilha = openpyxl.load_workbook(arquivo, read_only=True)
    aba = next(n for n in planilha.sheetnames if n.startswith("COVERAGE"))
    linhas = planilha[aba].iter_rows(values_only=True)
    cabecalho = next(linhas)
    col = cabecalho.index("geocode")
    return pd.DataFrame([l for l in linhas if str(l[col]) == CODIGO_IBGE], columns=cabecalho)


def areas_protegidas():
    pasta = os.path.join(PASTA_FONTES, "AreasProtegidas")
    os.makedirs(pasta, exist_ok=True)
    municipios = gpd.read_file(glob.glob(os.path.join(PASTA_FONTES, "PA_Municipios_2025", "*.shp"))[0])
    oriximina = municipios[municipios["CD_MUN"] == CODIGO_IBGE].to_crs(ALBERS).union_all()

    def recorta(g):
        g = g.to_crs(ALBERS)
        g = g[g.intersects(oriximina)].copy()
        g["geometry"] = g.intersection(oriximina)
        g["km2"] = g.area / 1e6
        return g[g["km2"] > 0.5]   # menos de meio km² é só a borda encostando no limite do município

    tis = recorta(gpd.read_file(baixar(URL_TIS, os.path.join(pasta, "terras_indigenas_regiao.json"))))
    arq_q = baixar(URL_QUILOMBOLAS, os.path.join(pasta, "quilombolas_incra.zip"))
    quilombolas = recorta(gpd.read_file(f"zip://{arq_q}"))
    with urllib.request.urlopen(urllib.request.Request(API_CNUC, headers={"User-Agent": "Mozilla/5.0"}),
                                timeout=60) as resposta:
        recursos = json.loads(resposta.read())["result"]["resources"]
    url_cnuc = max((r for r in recursos if r["url"].lower().endswith(".zip") and "shp_cnuc" in r["url"].lower()),
                   key=lambda r: r.get("created") or "")["url"]
    arq_uc = baixar(url_cnuc, os.path.join(pasta, os.path.basename(url_cnuc)))
    shp = next(n for n in zipfile.ZipFile(arq_uc).namelist() if n.endswith(".shp") and "pontos" not in n)
    ucs = gpd.read_file(f"zip://{arq_uc}!{shp}", encoding="latin1", bbox=LIMITES)
    for coluna in ("nome_uc", "categoria", "grupo", "esfera"):
        ucs[coluna] = ucs[coluna].map(sem_mojibake)
    ucs = recorta(ucs)
    return oriximina, tis, quilombolas, ucs


def main():
    # 1) Uso e cobertura da terra
    df = cobertura()
    anos = [c for c in df.columns if re.fullmatch(r"y\d{4}", c)]
    area = df.set_index("class")[anos] / 100   # hectares -> km²
    ultimo = anos[-1]
    df_uso = (pd.DataFrame({"Uso e cobertura": [CLASSES.get(c, str(c)) for c in area.index],
                            "Área (km²)": area[ultimo].round(1).values})
              .sort_values("Área (km²)", ascending=False))
    df_uso.to_csv(os.path.join(PASTA_COMPONENTES, "comp_310.csv"), sep=";", index=False, encoding="utf-8")
    hist = pd.DataFrame({"Ano": [int(a[1:]) for a in anos]})
    for classe in HISTORICO:
        hist[CLASSES[classe]] = area.loc[classe, anos].round(1).values if classe in area.index else 0
    hist.to_csv(os.path.join(PASTA_COMPONENTES, "comp_310_historico.csv"), sep=";", index=False, encoding="utf-8")
    print(df_uso.to_string(index=False))
    total = area[ultimo].sum()
    floresta = area.loc[[c for c in (3, 6, 4) if c in area.index], ultimo].sum()

    def em(classe, ano):
        return area.loc[classe, ano] if classe in area.index else 0

    # 2) Áreas protegidas
    oriximina, tis, quilombolas, ucs = areas_protegidas()
    # Uma linha por área (o tipo vai no nome: "TI", "Quilombo", "Estação Ecológica", "Floresta Nacional"...);
    # com uma coluna por tipo, cada linha usava só uma e as barras ficavam finas demais
    linhas = [(f"TI {nome_bonito(r['terrai_nome'])}", r["km2"]) for _, r in tis.iterrows()]
    linhas += [(f"Quilombo {nome_bonito(r['nm_comunid'])}", r["km2"]) for _, r in quilombolas.iterrows()]
    linhas += [(nome_bonito(r["nome_uc"]), r["km2"]) for _, r in ucs.iterrows()]
    df_prot = (pd.DataFrame(linhas, columns=["Área", "Área (km²)"]).sort_values("Área (km²)", ascending=False))
    df_prot["Área (km²)"] = df_prot["Área (km²)"].round().astype(int)
    df_prot.to_csv(os.path.join(PASTA_COMPONENTES, "comp_311.csv"), sep=";", index=False, encoding="utf-8")
    print(df_prot.to_string(index=False))

    def sem_repetir(*grupos):
        juntos = pd.concat([g.geometry for g in grupos])
        return gpd.GeoSeries([juntos.union_all()], crs=ALBERS).area.sum() / 1e6
    area_mun = oriximina.area / 1e6
    protegida = sem_repetir(tis, quilombolas, ucs)
    # Maior sobreposição entre áreas de tipos diferentes (para o exemplo da descrição)
    nomeadas = ([(f"a TI {nome_bonito(r['terrai_nome'])}", r.geometry) for _, r in tis.iterrows()]
                + [(f"o território quilombola {nome_bonito(r['nm_comunid'])}", r.geometry) for _, r in quilombolas.iterrows()])
    sobreposicoes = sorted(((a.intersection(r.geometry).area / 1e6, nome, nome_bonito(r["nome_uc"]))
                            for nome, a in nomeadas for _, r in ucs.iterrows()), reverse=True)
    maior = sobreposicoes[0]

    link_mb = {"nome": "MapBiomas – Estatísticas de cobertura e uso da terra", "url": "https://brasil.mapbiomas.org/estatisticas/"}
    componentes = {
        "comp_310": {
            "id": 310, "titulo": "Uso e cobertura da terra", "fonte": "MapBiomas | Coleção 11",
            "atualizacao": ultimo[1:], "icone": "grid_view", "area": "Município de Oriximiná",
            "resumo": f"Como são usados os {km2(total)} km² de Oriximiná: floresta, rios, campos, pastagem, mineração, cidade.",
            "descricao": (
                f"Mostra como estão cobertos os {km2(total)} km² de Oriximiná em {ultimo[1:]}, segundo o MapBiomas, "
                f"que mapeia o Brasil por imagens de satélite ano a ano desde 1985. A floresta ocupa "
                f"{km2(floresta)} km² ({floresta / total:.0%}); rios e lagos, {km2(em(33, ultimo))} km². A ocupação "
                f"humana é pequena, mas cresceu: a pastagem passou de {km2(em(15, anos[0]))} km² em {anos[0][1:]} para "
                f"{km2(em(15, ultimo))} km²; a mineração (a bauxita de Porto Trombetas), de {km2(em(30, anos[0]))} para "
                f"{km2(em(30, ultimo))} km²; a área urbana, de {em(24, anos[0]):.0f} para {em(24, ultimo):.0f} km². "
                "O Plano Diretor de Oriximiná tem um zoneamento, mas ele não está publicado como dados (o cartão "
                "do Taipei mostra o zoneamento da cidade); este cartão mostra o uso real da terra."),
            "exemplo": ("Ajuda a planejar o crescimento da cidade e a proteger a floresta: mostra onde a ocupação "
                        "avança (pastagem, mineração, área urbana) e quanto ainda é floresta e rio."),
            "tipos": ["mapa_arvore", "barras"],
            "cores": ["#2E7D32", "#56B96D", "#9DC56E", "#F8CF58", "#24B0DD", "#a0b8e8", "#F49F36", "#F65658",
                      "#c18ad8", "#8D6E63", "#bdbdbd", "#E57373", "#FFF176"],
            "unidade": "km²", "meta": None, "maximo": None, "links": [link_mb], "colaboradores": [], "mapa": None,
            "historico": {"tipo": "arquivo"},
        },
        "comp_311": {
            "id": 311, "titulo": "Áreas protegidas", "fonte": "FUNAI | INCRA | MMA – CNUC",
            "atualizacao": pd.Timestamp.now().strftime("%m/%Y"), "icone": "bar_chart", "area": "Município de Oriximiná",
            "resumo": "Terras indígenas, territórios quilombolas e unidades de conservação dentro de Oriximiná, em km².",
            "descricao": (
                f"Mostra quanto de cada área protegida fica dentro de Oriximiná: {len(tis)} terras indígenas "
                f"(FUNAI), {len(quilombolas)} territórios quilombolas (INCRA) e {len(ucs)} unidades de conservação "
                f"(Cadastro Nacional de Unidades de Conservação). Juntas, sem contar duas vezes as partes em que se "
                f"sobrepõem, cobrem {km2(protegida)} km² dos {km2(area_mun)} km² do município "
                f"({protegida / area_mun:.0%}). As áreas se sobrepõem em vários lugares (a maior sobreposição: "
                f"{maior[1]} e a {maior[2]} dividem cerca de {km2(maior[0])} km²), por isso a soma das áreas do "
                "gráfico passa da área protegida total. Em regra, a proteção integral (estação ecológica, reserva biológica) não "
                "permite moradia nem exploração (onde ela encontra territórios tradicionais, há acordos e disputas); o uso sustentável (florestas nacionais e estaduais) permite o uso controlado dos recursos. "
                "As áreas foram calculadas com o contorno do município do IBGE."),
            "exemplo": ("Mostra que quase todo o município é área protegida: o planejamento da cidade e das "
                        "comunidades precisa considerar as regras de cada área e os povos que vivem nelas."),
            "tipos": ["mapa_arvore", "barras"],
            "cores": ["#2E7D32", "#F49F36", "#9DC56E", "#56B96D", "#c18ad8", "#24B0DD", "#F8CF58", "#a0b8e8"],
            "unidade": "km²", "meta": None, "maximo": None,
            "links": [{"nome": "FUNAI – Terras indígenas (geoserver)", "url": "https://geoserver.funai.gov.br"},
                      {"nome": "INCRA – Acervo fundiário (quilombolas)", "url": "https://certificacao.incra.gov.br"},
                      {"nome": "MMA – Cadastro Nacional de Unidades de Conservação",
                       "url": "https://dados.mma.gov.br/dataset/unidadesdeconservacao"}],
            "colaboradores": [], "mapa": None, "historico": None,
        },
    }

    config = json.load(open(ARQ_ADMIN, encoding="utf-8"))
    config.setdefault("componentes_novos", {}).update(componentes)
    paineis = config["paineis_publicos"]
    planejamento = paineis.setdefault("planejamento", {"titulo": "Planejamento urbano", "icone": "home_work",
                                                       "cartoes": [], "grupo": "Painel de Oriximiná"})
    for chave in componentes:   # mantém a ordem escolhida; os que faltam entram no fim
        if chave not in planejamento["cartoes"]:
            planejamento["cartoes"].append(chave)
    with open(ARQ_ADMIN, "w", encoding="utf-8") as arquivo:
        json.dump(config, arquivo, ensure_ascii=False, indent=2)
    print(f"área protegida: {protegida:.0f} de {area_mun:.0f} km²; painel Planejamento urbano:", planejamento["cartoes"])


if __name__ == "__main__":
    main()
